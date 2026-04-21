#!/usr/bin/env python
# -*- coding: utf-8 -*-
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def parse_markdown_to_docx(md_file, docx_file):
    doc = Document()

    # 设置中文字体
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    doc.styles['Normal'].font.size = Pt(12)

    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i].rstrip()

        # 代码块处理
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                if code_lines:
                    p = doc.add_paragraph('\n'.join(code_lines))
                    p.style = 'Normal'
                    p.paragraph_format.left_indent = Inches(0.5)
                    for run in p.runs:
                        run.font.name = 'Consolas'
                        run.font.size = Pt(10)
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 表格处理
        if '|' in line and line.strip().startswith('|'):
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            # 检查下一行是否还是表格
            if i < len(lines) and '|' in lines[i] and lines[i].strip().startswith('|'):
                continue
            else:
                # 表格结束，创建表格
                in_table = False
                create_table(doc, table_lines)
                table_lines = []
            continue

        # 标题处理
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)
        # 引用块（截图占位）
        elif line.startswith('> '):
            p = doc.add_paragraph(line[2:])
            p.paragraph_format.left_indent = Inches(0.5)
            for run in p.runs:
                run.font.color.rgb = RGBColor(128, 128, 128)
                run.font.italic = True
        # 列表
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r'^\d+\.\s', line):
            doc.add_paragraph(re.sub(r'^\d+\.\s', '', line), style='List Number')
        # 空行
        elif line.strip() == '':
            pass
        # 普通段落
        else:
            # 处理粗体
            line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
            if line.strip():
                doc.add_paragraph(line)

        i += 1

    doc.save(docx_file)
    print(f'转换完成：{docx_file}')

def create_table(doc, table_lines):
    # 过滤掉分隔行
    data_lines = [line for line in table_lines if not re.match(r'^\|[\s\-:]+\|', line)]
    if not data_lines:
        return

    # 解析表格数据
    rows = []
    for line in data_lines:
        cells = [cell.strip() for cell in line.split('|')[1:-1]]
        rows.append(cells)

    if not rows:
        return

    # 创建表格
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Light Grid Accent 1'

    for i, row_data in enumerate(rows):
        for j, cell_data in enumerate(row_data):
            table.rows[i].cells[j].text = cell_data

if __name__ == '__main__':
    parse_markdown_to_docx(
        'E:/A毕业设计/plnm/system/thesis_revised.md',
        'E:/A毕业设计/plnm/system/基于智能体协作的学习生活一体化管理系统_修订版.docx'
    )
