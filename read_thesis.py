#!/usr/bin/env python
# -*- coding: utf-8 -*-
import sys
import io
import docx
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

thesis_path = Path("E:/A毕业设计/plnm/system/基于智能体协作的学习生活一体化管理系统.docx")

with open(thesis_path, 'rb') as f:
    doc = docx.Document(io.BytesIO(f.read()))

print(f'总段落数: {len(doc.paragraphs)}\n')

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        print(f'{text}')
